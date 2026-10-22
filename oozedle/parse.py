"""Read both generations of tournament workbooks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

STAT_LABELS = {
    "Points Played": "points_played",
    "Assists": "assists",
    "Goals": "goals",
    "Blocks": "blocks",
    "D's": "blocks",
    "Turnovers": "turnovers",
    "Turns": "turnovers",
}
STATS = ("points_played", "assists", "goals", "blocks", "turnovers")


@dataclass
class Game:
    index: int
    opponent_label: str
    opponent: str
    goals_for: float
    goals_against: float
    team_points: int | None
    team_blocks: int | None
    team_turnovers: int | None
    clean_holds: int | None


@dataclass
class PlayerLine:
    name: str
    points_played: int
    assists: int
    goals: int
    blocks: int
    turnovers: int


@dataclass
class Workbook:
    summary: list[PlayerLine]
    games: list[Game]
    player_games: dict[int, list[PlayerLine]]
    layout: str
    warnings: list[str]


def normalize_opponent(label: str) -> str:
    """Remove a game-number suffix while preserving the opponent's spelling."""
    return re.sub(r"(?<=[^\W\d_])\d+$", "", label.strip())


def _cell(rows: list[tuple], r: int, c: int):
    return rows[r][c] if 0 <= r < len(rows) and 0 <= c < len(rows[r]) else None


def _line(
    name: object, values: dict, warnings: list[str], summary: bool = False
) -> PlayerLine | None:
    name = str(name or "").strip()
    if not name or name in {"Total", "Total/Pts"}:
        return None
    stats = {}
    for key in STATS:
        value = values.get(key)
        try:
            stats[key] = int(value or 0)
        except (TypeError, ValueError, OverflowError):
            warnings.append(f"{name}: nonnumeric {key} {value!r}; row skipped")
            return None
    keep = (
        any(v is not None for v in values.values())
        if summary
        else any(v > 0 for v in stats.values())
    )
    return PlayerLine(name, **stats) if keep else None


def parse_workbook(path: Path) -> Workbook:
    """Parse Summary totals and positionally matched per-game lines."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {sheet.title: list(sheet.values) for sheet in book.worksheets}
    finally:
        book.close()
    rows = sheets["Summary"]
    warnings: list[str] = []
    header = next(i for i, row in enumerate(rows) if "Points Played" in row)
    columns = {STAT_LABELS[v]: c for c, v in enumerate(rows[header]) if v in STAT_LABELS}
    summary = []
    for row in rows[header + 1 :]:
        if not row[0]:
            break
        line = _line(row[0], {key: row[c] for key, c in columns.items()}, warnings, summary=True)
        if line:
            summary.append(line)
    position = next(
        (
            (r, c)
            for r, row in enumerate(rows)
            for c, value in enumerate(row)
            if value == "Goals Scored"
        ),
        None,
    )
    r, c = position if position is not None else (0, len(rows[0]))
    labels = []
    for col in range(c + 1, len(rows[r - 1])):
        value = _cell(rows, r - 1, col)
        if value is None or value == "":
            break
        labels.append(str(value).strip())
    if not labels and isinstance(_cell(rows, r, c + 1), (int, float)):
        labels = ["Alumni"]
    game_rows = {
        _cell(rows, row, c): row
        for row in range(r, len(rows))
        if isinstance(_cell(rows, row, c), str)
    }

    def value(label: str, col: int, optional: bool = False):
        row = game_rows.get(label)
        val = _cell(rows, row, col) if row is not None else None
        return None if optional and val is None else (val or 0)

    games = [
        Game(
            index=i,
            opponent_label=label,
            opponent=normalize_opponent(label),
            goals_for=float(value("Goals Scored", c + i)),
            goals_against=float(value("Goals Given", c + i)),
            team_points=_optional_int(value("Total Points", c + i, True)),
            team_blocks=_optional_int(
                value("Blocks" if "Blocks" in game_rows else "D's", c + i, True)
            ),
            team_turnovers=_optional_int(value("Turnovers", c + i, True)),
            clean_holds=_optional_int(value("Clean Holds", c + i, True)),
        )
        for i, label in enumerate(labels, 1)
    ]
    game_sheets = [
        (name, data)
        for name, data in sheets.items()
        if name != "Summary" and _cell(data, 0, 0) == "Name"
    ]
    layout = "B" if game_sheets else "A"
    player_games: dict[int, list[PlayerLine]] = {}
    usable = True

    def check_label(actual: str, index: int):
        expected = labels[index - 1]
        if normalize_opponent(actual).casefold() != normalize_opponent(expected).casefold():
            warnings.append(f"Game {index} label mismatch: {actual!r} versus {expected!r}")

    if layout == "B":
        if len(game_sheets) != len(games):
            usable = False
        else:
            for i, (name, data) in enumerate(game_sheets, 1):
                check_label(name, i)
                cols = {STAT_LABELS[v]: j for j, v in enumerate(data[0]) if v in STAT_LABELS}
                player_games[i] = []
                for row in data[1:]:
                    line = _line(row[0], {key: row[j] for key, j in cols.items()}, warnings)
                    if line:
                        player_games[i].append(line)
    else:
        for name, data in sheets.items():
            if name == "Summary" or _cell(data, 0, 0) not in (None, ""):
                continue
            stat_rows = {STAT_LABELS[row[0]]: row for row in data if row[0] in STAT_LABELS}
            if "points_played" not in stat_rows:
                continue
            game_cols = []
            for col, label in enumerate(data[0][1:], 1):
                if label in {"Total", "Total/Pts"} or label is None:
                    break
                game_cols.append((col, str(label)))
            if len(game_cols) != len(games):
                usable = False
                break
            for i, (col, label) in enumerate(game_cols, 1):
                check_label(label, i)
                line = _line(name, {key: row[col] for key, row in stat_rows.items()}, warnings)
                if line:
                    player_games.setdefault(i, []).append(line)
    if not usable:
        warnings.append("Per-game count differs from Summary; dropped all per-game data")
        player_games = {}
    if len(games) == 1 and not player_games:
        player_games = {
            1: [line for line in summary if any(getattr(line, stat) > 0 for stat in STATS)]
        }
        warnings.append("Synthesized single-game player lines from Summary totals")
    return Workbook(summary, games, player_games, layout, list(dict.fromkeys(warnings)))


def _optional_int(value) -> int | None:
    return None if value is None else int(value)
