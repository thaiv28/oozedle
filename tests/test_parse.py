from __future__ import annotations

from conftest import make_layout_a, make_single_game_with_stale_sheets
from openpyxl import load_workbook

from oozedle.parse import parse_workbook


def test_both_layouts(tmp_path, workbook_factory):
    workbook = parse_workbook(workbook_factory(tmp_path / "stats.xlsx"))
    assert [g.opponent_label for g in workbook.games] == ["SJSU1", "Chico", "Chico"]
    assert [g.opponent for g in workbook.games] == ["SJSU", "Chico", "Chico"]
    assert workbook.games[0].goals_for == 15.5
    assert [line.name for line in workbook.summary] == ["Ada", "Ben"]
    assert [line.points_played for line in workbook.player_games[1]] == [4, 6]
    assert [line.name for line in workbook.player_games[2]] == ["Ada"]
    assert workbook.player_games[2][0].blocks == 1
    assert workbook.player_games[3][0].blocks == 0
    assert not workbook.warnings


def test_shifted_columns(tmp_path, workbook_factory):
    workbook = parse_workbook(workbook_factory(tmp_path / "stats.xlsx", shifted=True))
    assert workbook.summary[0].goals == 3
    assert workbook.summary[0].blocks == 1
    assert len(workbook.games) == 3


def test_single_game_stale(tmp_path):
    workbook = parse_workbook(make_single_game_with_stale_sheets(tmp_path / "alumni.xlsx"))
    assert len(workbook.games) == 1
    assert workbook.games[0].opponent == "Alumni"
    assert workbook.player_games[1] == workbook.summary
    assert any("count differs" in w for w in workbook.warnings)
    assert any("Synthesized" in w for w in workbook.warnings)


def test_label_mismatch(tmp_path):
    path = make_layout_a(tmp_path / "stats.xlsx")
    book = load_workbook(path)
    book["Ada"].cell(1, 2, "UCI B")
    book.save(path)
    workbook = parse_workbook(path)
    assert any("label mismatch" in w for w in workbook.warnings)
    assert workbook.player_games[1][0].points_played == 4


def test_a_count_mismatch_discards_all_players(tmp_path):
    path = make_layout_a(tmp_path / "stats.xlsx")
    book = load_workbook(path)
    book["Ben"].cell(1, 3, "Total")
    book.save(path)
    workbook = parse_workbook(path)
    assert workbook.player_games == {}
    assert any("count differs" in w for w in workbook.warnings)


def test_zero_summary_and_no_games(tmp_path):
    path = make_layout_a(tmp_path / "stats.xlsx")
    book = load_workbook(path)
    sheet = book["Summary"]
    for column in range(2, 7):
        sheet.cell(3, column, 0)
    for row in sheet:
        for cell in row:
            if cell.column >= 12:
                cell.value = None
    book.save(path)
    workbook = parse_workbook(path)
    assert workbook.summary[0].name == "Ada"
    assert workbook.games == []


def test_nonnumeric_annotation_skips_row(tmp_path):
    path = make_layout_a(tmp_path / "stats.xlsx")
    book = load_workbook(path)
    first = book["Summary"].cell(3, 1).value
    book["Summary"].cell(3, 3, "Bros locked up: 1")
    book.save(path)
    workbook = parse_workbook(path)
    assert first not in [line.name for line in workbook.summary]
    assert any("nonnumeric assists" in w and "row skipped" in w for w in workbook.warnings)
