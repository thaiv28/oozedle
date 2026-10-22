"""Small offline workbooks representing the two source layouts."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook


def _summary(book, labels, shifted=False, single=False):
    sheet = book.active
    sheet.title = "Summary"
    headers = ["Name", "Points Played", "Assists", "Goals", "D's", "Turnovers"]
    for col, value in enumerate(headers, 1):
        sheet.cell(2, col, value)
    for row, name, values in [
        (3, "Ada", [12.0, 2, 3, 1, 2]),
        (4, "Ben", [6, 1, 1, 0, 1]),
        (5, "DNP", [None] * 5),
    ]:
        sheet.cell(row, 1, name)
        for label, value in zip(["Points Played", "Assists", "Goals", "D's", "Turnovers"], values):
            sheet.cell(row, headers.index(label) + 1, value)
    col = 13 if shifted else 12
    if single:
        sheet.cell(1, col, "Game Stats")
        start = 2
    else:
        sheet.cell(1, col, "Game Stats")
        for i, label in enumerate(labels, col + 1):
            sheet.cell(2, i, label)
        start = 3
    for offset, label in enumerate(
        ["Goals Scored", "Goals Given", "Total Points", "D's", "Turnovers", "Clean Holds"]
    ):
        sheet.cell(start + offset, col, label)
        for i in range(len(labels)):
            values = [15.5 if i == 0 else 8, 10, 25, 2, 3, 4]
            sheet.cell(start + offset, col + i + 1, values[offset])


def make_layout_a(path: Path, labels=("SJSU1", "Chico", "Chico"), shifted=False):
    book = Workbook()
    _summary(book, labels, shifted)
    stats = ["Points Played", "Assists", "Goals", "D's", "Turnovers"]
    for name, values in [
        ("Ada", [[4, 4, 4], [1, 1, 0], [1, 1, 1], [0, 1, 0], [1, 0, 1]]),
        (
            "Ben",
            [[6, None, None], [1, None, None], [1, None, None], [0, None, None], [1, None, None]],
        ),
    ]:
        sheet = book.create_sheet(name)
        sheet.append([None, *labels, "Total", "Total/Pts"])
        for label, row in zip(stats, values):
            sheet.append([label, *row[: len(labels)], sum(v or 0 for v in row)])
    book.save(path)
    return path


def make_layout_b(path: Path, labels=("SJSU1", "Chico", "Chico"), shifted=False):
    book = Workbook()
    _summary(book, labels, shifted)
    for i, label in enumerate(labels):
        sheet = book.create_sheet(label if label not in book.sheetnames else f"{label}2")
        sheet.append(["Name", "Points Played", "Assists", "Goals", "Blocks", "Turns"])
        sheet.append(["Ada", 4, 1 if i < 2 else 0, 1, 1 if i == 1 else 0, 0 if i == 1 else 1])
        if i == 0:
            sheet.append(["Ben", 6, 1, 1, 0, 1])
        sheet.append(["DNP", None, None, None, None, None])
        sheet.append(["Total", 10, 2, 2, 0, 2])
    sheet = book.create_sheet("Ada")
    sheet.append([None, "Points Played", "Assists", "Goals", "Blocks", "Turns"])
    book.save(path)
    return path


def make_single_game_with_stale_sheets(path: Path):
    make_layout_b(path)
    from openpyxl import load_workbook

    book = load_workbook(path)
    del book["Summary"]
    book.create_sheet("Summary", 0)
    book.active = 0
    _summary(book, ["Alumni"], single=True)
    book.save(path)
    return path


@pytest.fixture(params=[make_layout_a, make_layout_b], ids=["A", "B"])
def workbook_factory(request):
    return request.param
