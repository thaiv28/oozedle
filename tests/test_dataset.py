from __future__ import annotations

from conftest import make_layout_a
from openpyxl import load_workbook

from oozedle.dataset import build_dataset, load_aliases, slugify


def test_tables_order_aliases_validation_csv(tmp_path):
    season = tmp_path / "2024"
    season.mkdir()
    paths = [make_layout_a(season / f"{name} Stats.xlsx") for name in ["Zebra", "Alpha", "Last"]]
    manifest = [
        dict(season=2024, name=p.stem, path=f"2024/{p.name}", created=date, modified=None)
        for p, date in zip(paths[:2], ["2024-01-01T00:00:00Z", "2024-02-01T00:00:00Z"])
    ]
    book = load_workbook(paths[0])
    book["Ada"].cell(3, 2, 4)
    book.save(paths[0])
    dataset = build_dataset(tmp_path, manifest, {"Ada": "A. Player", "Ben": "A. Player"})
    assert dataset.tournaments.name.tolist() == ["Zebra", "Alpha", "Last"]
    assert dataset.tournaments.order.tolist() == [1, 2, 3]
    assert dataset.tournaments.tournament_id.tolist() == ["2024-zebra", "2024-alpha", "2024-last"]
    assert dataset.tournaments.wins.tolist() == [1, 1, 1]
    assert dataset.tournaments.losses.tolist() == [2, 2, 2]
    assert dataset.players.player_id.tolist() == ["a-player"]
    assert dataset.players.iloc[0].tournaments == 3
    line = dataset.player_tournaments.iloc[0]
    assert (line.games_played, line.wins, line.losses, line.points_played) == (3, 1, 2, 18)
    assert any("merged aliases" in w for w in dataset.warnings)
    assert any("assists" in w and "differ" in w for w in dataset.warnings)
    assert any("missing from manifest" in w for w in dataset.warnings)
    assert str(dataset.games.team_points.dtype) == "Int64"
    output = tmp_path / "csv"
    dataset.write_csv(output)
    assert len(list(output.glob("*.csv"))) == 5


def test_empty_and_aliases(tmp_path):
    dataset = build_dataset(tmp_path, [], {})
    assert dataset.games.empty
    assert dataset.players.columns.tolist()[0] == "player_id"
    assert load_aliases(tmp_path / "missing.yml") == {}
    path = tmp_path / "players.yml"
    path.write_text("aliases: {Nickname: Full Name}\n")
    assert load_aliases(path) == {"Nickname": "Full Name"}
    assert slugify("Héllo, World!") == "hello-world"


def test_unlisted_sorts_after_manifest_without_created(tmp_path):
    season = tmp_path / "2024"
    season.mkdir()
    for name in ["Z", "A"]:
        make_layout_a(season / f"{name} Stats.xlsx")
    manifest = [dict(season=2024, name="Z Stats", path="2024/Z Stats.xlsx", created=None)]
    dataset = build_dataset(tmp_path, manifest, {})
    assert dataset.tournaments.name.tolist() == ["Z", "A"]


def test_order_override_puts_listed_tournaments_in_order(tmp_path):
    from oozedle.dataset import load_order, tournament_name

    path = tmp_path / "tournaments.yml"
    path.write_text("order:\n  2026: [Regionals, Alumni Game]\n")
    assert load_order(path) == {2026: ["Regionals", "Alumni Game"]}
    assert load_order(tmp_path / "missing.yml") == {}
    assert tournament_name({"name": "Alumni Game Stats"}) == "Alumni Game"
